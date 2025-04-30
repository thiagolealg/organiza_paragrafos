import argparse
from docx import Document
from sentence_transformers import SentenceTransformer
import numpy as np
from scipy.optimize import linear_sum_assignment
from tqdm import tqdm # Adicionado para barra de progresso
import sys

def parse_args():
    ap = argparse.ArgumentParser(
        description="Reordena parágrafos de um .docx usando similaridade semântica.",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter # Melhora a ajuda
    )
    ap.add_argument("input_docx",  help="Caminho do arquivo Word (.docx) a ser reordenado.")
    ap.add_argument("-r", "--reference", default="ref_order.txt",
                    help="Arquivo texto (.txt) contendo os parágrafos na ordem correta, um por linha.")
    ap.add_argument("-o", "--output", default="sorted_output.docx",
                    help="Nome do arquivo .docx de saída com os parágrafos reordenados.")
    ap.add_argument("-m", "--model",
                    default="sentence-transformers/all-MiniLM-L6-v2",
                    help="Nome ou caminho do modelo sentence-transformers a ser usado.")
    ap.add_argument("-e", "--encoding", default="utf-8",
                    help="Encoding do arquivo de referência.") # Adicionado encoding
    return ap.parse_args()

def load_docx_paragraphs(path):
    """Carrega parágrafos não vazios de um arquivo .docx."""
    try:
        doc = Document(path)
        return [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    except Exception as e:
        print(f"Erro ao ler o arquivo DOCX '{path}': {e}", file=sys.stderr)
        sys.exit(1)

def load_reference_paragraphs(path, encoding='utf-8'):
    """Carrega parágrafos não vazios de um arquivo de texto."""
    try:
        with open(path, 'r', encoding=encoding) as f:
            return [line.strip() for line in f if line.strip()]
    except FileNotFoundError:
        print(f"Erro: Arquivo de referência '{path}' não encontrado.", file=sys.stderr)
        sys.exit(1)
    except Exception as e:
        print(f"Erro ao ler o arquivo de referência '{path}': {e}", file=sys.stderr)
        sys.exit(1)

def match_order(candidate_paragraphs, reference_paragraphs, model_name):
    """Calcula a similaridade e encontra a melhor ordem correspondente."""
    if not candidate_paragraphs:
        print("Aviso: Nenhum parágrafo encontrado no arquivo DOCX de entrada.", file=sys.stderr)
        return []
    if not reference_paragraphs:
        print("Erro: Nenhum parágrafo encontrado no arquivo de referência.", file=sys.stderr)
        sys.exit(1)

    if len(candidate_paragraphs) != len(reference_paragraphs):
        print(
            f"Aviso: O número de parágrafos difere entre a entrada ({len(candidate_paragraphs)}) "
            f"e a referência ({len(reference_paragraphs)}). O resultado pode não ser o esperado.",
            file=sys.stderr
        )
        # Poderia truncar/ajustar aqui, mas por ora só avisa

    print(f"Carregando modelo de embedding '{model_name}'...")
    try:
        model = SentenceTransformer(model_name)
    except Exception as e:
        print(f"Erro ao carregar o modelo '{model_name}': {e}", file=sys.stderr)
        print("Verifique se o nome está correto e se há conexão com a internet (para download inicial).", file=sys.stderr)
        sys.exit(1)

    print("Calculando embeddings para os parágrafos de entrada...")
    cand_emb = model.encode(candidate_paragraphs, show_progress_bar=True, normalize_embeddings=True)

    print("Calculando embeddings para os parágrafos de referência...")
    ref_emb = model.encode(reference_paragraphs, show_progress_bar=True, normalize_embeddings=True)

    # Calcula a matriz de similaridade cosseno
    cosine_sim = np.matmul(ref_emb, cand_emb.T)

    # Converte similaridade para 'custo' (distância) para o algoritmo húngaro
    # Maximizar similaridade é minimizar (1 - similaridade)
    cost_matrix = 1 - cosine_sim

    print("Encontrando a melhor correspondência de ordem (Algoritmo Húngaro)...")
    # Aplica o algoritmo de atribuição linear (Húngaro)
    # row_ind[i] (índice da referência) deve corresponder a col_ind[i] (índice do candidato)
    row_ind, col_ind = linear_sum_assignment(cost_matrix)

    # Cria um dicionário para mapear o índice original do candidato para o parágrafo
    original_candidate_map = {i: p for i, p in enumerate(candidate_paragraphs)}

    # Reconstrói a lista ordenada baseado na correspondência encontrada
    # O resultado `col_ind` já nos dá os índices dos *candidatos* na ordem que
    # melhor corresponde aos *referência* (que estão ordenados de 0 a N-1 implicitamente
    # pela ordem em `ref_emb` e `row_ind`).
    ordered_paragraphs = [original_candidate_map[col_idx] for col_idx in col_ind]

    # Caso haja mais candidatos que referências (ou vice-versa após o aviso),
    # precisamos adicionar os parágrafos 'extras' que não foram mapeados.
    # Esta parte pode precisar de lógica mais sofisticada dependendo do caso de uso.
    if len(ordered_paragraphs) < len(candidate_paragraphs):
        print("Aviso: Alguns parágrafos de entrada não foram mapeados para a referência.", file=sys.stderr)
        # Simplesmente adiciona os não mapeados ao final (pode não ser o ideal)
        mapped_indices = set(col_ind)
        for i, p in enumerate(candidate_paragraphs):
            if i not in mapped_indices:
                ordered_paragraphs.append(p)

    return ordered_paragraphs

def write_docx(paragraphs, path):
    """Escreve uma lista de strings como parágrafos em um novo arquivo .docx."""
    doc = Document()
    print(f"Escrevendo {len(paragraphs)} parágrafos no arquivo de saída '{path}'...")
    for text in paragraphs:
        doc.add_paragraph(text)
    try:
        doc.save(path)
    except Exception as e:
        print(f"Erro ao salvar o arquivo DOCX '{path}': {e}", file=sys.stderr)
        sys.exit(1)

def main():
    args = parse_args()

    print(f"Lendo parágrafos do arquivo de entrada: {args.input_docx}")
    candidate_paragraphs = load_docx_paragraphs(args.input_docx)

    print(f"Lendo parágrafos de referência: {args.reference} (encoding: {args.encoding})")
    reference_paragraphs = load_reference_paragraphs(args.reference, args.encoding)

    ordered_paragraphs = match_order(candidate_paragraphs, reference_paragraphs, args.model)

    if ordered_paragraphs:
        write_docx(ordered_paragraphs, args.output)
        print(f"\nProcesso concluído. Arquivo reordenado salvo em: {args.output}")
    else:
        print("\nNenhum parágrafo para escrever no arquivo de saída.")

if __name__ == "__main__":
    main()
